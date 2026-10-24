import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

doc = docx.Document()

# Set standard 0.75 in margins
for sec in doc.sections:
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)

# Styling colors
NAVY = RGBColor(30, 58, 138)       # #1E3A8A
SLATE = RGBColor(30, 41, 59)       # #1E293B
TEAL = RGBColor(2, 132, 199)       # #0284C7
GRAY = RGBColor(100, 116, 139)     # #64748B
GREEN = RGBColor(22, 163, 74)      # #16A34A

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def add_header_p(text, level=1):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = NAVY
        run.font.bold = True
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = TEAL
        run.font.bold = True
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.size = Pt(11)
        run.font.color.rgb = SLATE
        run.font.bold = True
        h.paragraph_format.space_before = Pt(6)
        h.paragraph_format.space_after = Pt(2)
    return h

def add_callout(text, title='FACULTY DEMONSTRATION VERIFICATION'):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, 'F0FDF4')
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    r_title = p.add_run(f'✓ {title}\n')
    r_title.font.bold = True
    r_title.font.size = Pt(11)
    r_title.font.color.rgb = GREEN
    r_text = p.add_run(text)
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = SLATE
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_figure(path, caption, width=Inches(5.5)):
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(path, width=width)
        
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(2)
        cap.paragraph_format.space_after = Pt(8)
        run_cap = cap.add_run(caption)
        run_cap.font.size = Pt(9)
        run_cap.font.italic = True
        run_cap.font.color.rgb = GRAY

# ----------------- COVER / HEADER -----------------
title_p = doc.add_paragraph()
title_p.paragraph_format.space_before = Pt(10)
title_p.paragraph_format.space_after = Pt(2)
r_title = title_p.add_run('COG: Autonomous Floor Waste-Detection Robot')
r_title.font.size = Pt(22)
r_title.font.bold = True
r_title.font.color.rgb = NAVY

sub_p = doc.add_paragraph()
sub_p.paragraph_format.space_before = Pt(0)
sub_p.paragraph_format.space_after = Pt(12)
r_sub = sub_p.add_run('Closed-Loop Visual Servoing, Edge YOLOv8 Classification & Embedded 4WD Control\nDA3 Course Project Submission & Demonstration Report')
r_sub.font.size = Pt(12)
r_sub.font.color.rgb = TEAL

# Metadata Table
meta_tbl = doc.add_table(rows=4, cols=2)
meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_data = [
    ('Assessment:', 'DA3 Submission: Codes, Readme, Flowchart, Working Screen & Output'),
    ('Due Date:', 'November 3, 2026 11:59 PM'),
    ('Demonstration Status:', 'Successfully Demonstrated Live Before Course Faculty (Full Working Confirmed)'),
    ('Hardware Setup:', 'Raspberry Pi 4B, Pi Camera V2, Arduino Uno, L298N Dual H-Bridge, 4WD Chassis')
]
for i, (k, v) in enumerate(meta_data):
    row = meta_tbl.rows[i]
    c0 = row.cells[0]
    c1 = row.cells[1]
    c0.width = Inches(1.8)
    c1.width = Inches(4.8)
    set_cell_background(c0, 'F1F5F9')
    set_cell_background(c1, 'F8FAFC')
    p0 = c0.paragraphs[0]
    p0.add_run(k).font.bold = True
    p0.runs[0].font.size = Pt(9.5)
    p0.runs[0].font.color.rgb = SLATE
    p1 = c1.paragraphs[0]
    p1.add_run(v).font.size = Pt(9.5)
    p1.runs[0].font.color.rgb = SLATE

doc.add_paragraph().paragraph_format.space_after = Pt(6)

add_callout(
    'The physical 4WD autonomous robot car was evaluated in a live floor laboratory demonstration before the course faculty. '
    'The robot successfully performed autonomous 360° sweeping, detected a floor target (paper cup), aligned its heading through closed-loop visual servoing, '
    'pulsed forward to approach the target, reached the arrival threshold, classified the item as WASTE, illuminated the Green indicator LED for 3.0 seconds, '
    'and executed an obstacle-clearing reverse and pivot maneuver to continue searching for other items. The faculty confirmed full autonomous operation.',
    'FACULTY PRESENTATION SUCCESS & VERIFICATION'
)

# ----------------- SECTION 1: SYSTEM ARCHITECTURE -----------------
add_header_p('1. System Architecture & Dual-Tier Computing', level=1)
p_arch = doc.add_paragraph(
    'The COG robot employs a dual-tier edge compute architecture designed to separate high-level computer vision and state management '
    'from hard real-time motor actuation and distance sensing. The Raspberry Pi 4B runs an optimized YOLOv8 neural network and the master autonomous '
    'state machine, communicating with an Arduino Uno over a high-speed non-blocking USB serial protocol at 115200 baud.'
)
p_arch.runs[0].font.size = Pt(10)

# Architecture table
arch_tbl = doc.add_table(rows=6, cols=3)
arch_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ['Subsystem', 'Hardware Component', 'Primary Responsibility']
for j, h in enumerate(headers):
    c = arch_tbl.cell(0, j)
    set_cell_background(c, '1E3A8A')
    p = c.paragraphs[0]
    r = p.add_run(h)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    r.font.size = Pt(9.5)

arch_rows = [
    ('High-Level Compute', 'Raspberry Pi 4B (4GB, Quad Cortex-A72)', 'YOLOv8n object detection (224x224), 7-state finite state machine, decision logic, LED signaling.'),
    ('Vision Sensor', 'Raspberry Pi Camera Module V2 (8MP Sony IMX219)', 'Floor-angled optical capture, live video streaming, target bearing angle calculation.'),
    ('Low-Level Microcontroller', 'Arduino Uno (ATmega328P)', 'Hardware PWM generation for L298N motor driver, non-blocking timed pulses, HC-SR04 ultrasonic echo polling.'),
    ('Motor Actuation', 'L298N Dual H-Bridge + 4x Geared DC Motors', 'Pulsed skid-steer navigation: Forward (F), Backward (B), Pivot Left (TL), Pivot Right (TR), Stop (S).'),
    ('Telemetry & Obstacle', 'HC-SR04 Ultrasonic Distance Sensor', 'High-frequency distance measurement (2cm to 400cm) for arrival and collision avoidance.')
]
for i, row_data in enumerate(arch_rows):
    row = arch_tbl.rows[i+1]
    bg = 'FFFFFF' if i % 2 == 0 else 'F8FAFC'
    for j, val in enumerate(row_data):
        c = row.cells[j]
        set_cell_background(c, bg)
        p = c.paragraphs[0]
        r = p.add_run(val)
        r.font.size = Pt(9)
        r.font.color.rgb = SLATE
        if j == 0:
            r.font.bold = True

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# ----------------- SECTION 2: FLOWCHART -----------------
add_header_p('2. System Flowchart & Finite State Machine', level=1)
p_fc = doc.add_paragraph(
    'The robot is driven by an autonomous 7-state finite state machine (FSM). Visual servoing provides closed-loop proportional control during alignment, '
    'while ballistic forward pulsing safely approaches target objects on uneven laboratory flooring without motor over-travel.'
)
p_fc.runs[0].font.size = Pt(10)

add_figure('diagrams/system_flowchart_light.png', 'Figure 1: Complete System Flowchart and Autonomous 7-State FSM Execution Pipeline', width=Inches(6.2))

# State explanation table
fsm_tbl = doc.add_table(rows=8, cols=3)
fsm_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
headers_fsm = ['State', 'Trigger / Condition', 'Execution Behavior & Serial Command']
for j, h in enumerate(headers_fsm):
    c = fsm_tbl.cell(0, j)
    set_cell_background(c, '0284C7')
    p = c.paragraphs[0]
    r = p.add_run(h)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    r.font.size = Pt(9.5)

fsm_rows = [
    ('1. SCAN', 'Sweep start / target lost', 'Executes incremental 20° pivot pulse (TR,80). Captures frame, runs YOLOv8. Sweeps full 360° until object detected.'),
    ('2. ALIGN', 'Target detected in frame', 'Computes optical bearing error θ = (x_center - W/2) * (HFOV / W). If |θ| > 3°, pulses pivot TL or TR until centered.'),
    ('3. APPROACH', 'Bearing aligned (|θ| <= 3°)', 'Drives forward in controlled pulses (F,300). Continuously checks arrival criteria (bbox bottom >= 92% of frame height or ultrasonic <= 12cm).'),
    ('4. VERDICT', 'Arrival threshold met', 'Robot stops (S). Collects 3 consecutive inference frames, calculates mean confidence, and looks up class in Waste vs Non-Waste rule tables.'),
    ('5. INDICATE', 'Classification established', 'If Waste: turns ON Green LED (BCM 27). If Non-Waste: turns ON Red LED (BCM 17). Holds status illumination for 3.0 seconds.'),
    ('6. LEAVE', 'Indication elapsed', 'Backs away from the inspected item (B,500) and executes large escape pivot (TR,960) to avoid re-detecting the same object. Resumes SCAN.'),
    ('7. DONE', 'All sweeps completed', 'Blinks indicator LEDs 3 times, parks motors, closes serial connection, and issues safe operating system poweroff.')
]
for i, row_data in enumerate(fsm_rows):
    row = fsm_tbl.rows[i+1]
    bg = 'FFFFFF' if i % 2 == 0 else 'F8FAFC'
    for j, val in enumerate(row_data):
        c = row.cells[j]
        set_cell_background(c, bg)
        p = c.paragraphs[0]
        r = p.add_run(val)
        r.font.size = Pt(9)
        r.font.color.rgb = SLATE
        if j == 0:
            r.font.bold = True

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# ----------------- SECTION 3: WORKING SCREEN SCREENSHOTS -----------------
add_header_p('3. Working Screen Screenshots & Real-Time Output', level=1)
p_ws = doc.add_paragraph(
    'The following screenshots demonstrate the real-time working screen of the Raspberry Pi during the live testing and presentation. '
    'The system displays an interactive Heads-Up Display (HUD) showing live bounding boxes, object confidence, bearing alignment guides, '
    'and the active FSM state alongside real-time terminal telemetry.'
)
p_ws.runs[0].font.size = Pt(10)

add_figure('docx_figures/hud_preview_closeup.jpg', 'Figure 2: Real-time Camera HUD Preview during APPROACH State with Bounding Box [cup 0.78]', width=Inches(3.8))
add_figure('docx_figures/Screenshot_2026-10-22-11-20-47-63_cropped.jpg', 'Figure 3: Working Screen Desktop View — Real-time Terminal Log (Left) and Camera Preview HUD (Right) at Arrival & Verdict', width=Inches(6.0))
add_figure('docx_figures/terminal_log_closeup.jpg', 'Figure 4: Terminal Telemetry Closeup Showing Scan Sweep, Bearing Alignment, Approach Pulses, and Green LED Waste Verdict', width=Inches(3.8))
add_figure('docx_figures/Screenshot_2026-10-22-11-20-50-85_cropped.jpg', 'Figure 5: Working Screen Showing Transition from INDICATE to LEAVE State (Reversing B,500 and Pivoting TR,960)', width=Inches(6.0))

# ----------------- SECTION 4: PHYSICAL ROBOT & DEMONSTRATION SETUP -----------------
add_header_p('4. Physical Robot Hardware & Floor Test Arena', level=1)
p_hw = doc.add_paragraph(
    'Photographic documentation of the physical 4WD robot car, sensor placement, and floor demonstration setup from the presentation session. '
    'The camera was angled toward the floor tiles to detect floor debris (such as paper cups and plastic bottles).'
)
p_hw.runs[0].font.size = Pt(10)

add_figure('docx_figures/floor_robot_setup.jpg', 'Figure 6: 4WD Robot Car Chassis on Floor Arena with Target Objects (Paper Cup and Plastic Water Bottle)', width=Inches(5.0))
add_figure('docx_figures/robot_approaching_cup.jpg', 'Figure 7: Physical Robot Performing Closed-Loop Navigation Toward the Paper Cup Target', width=Inches(5.0))

# ----------------- SECTION 5: LIVE OUTPUT & TELEMETRY LOGS -----------------
add_header_p('5. Live Demonstration Output & Serial Telemetry Log', level=1)
p_log_intro = doc.add_paragraph(
    'Below is the actual telemetry output recorded from the robot terminal during the successful live faculty evaluation:'
)
p_log_intro.runs[0].font.size = Pt(10)

# Terminal Log Box
log_table = doc.add_table(rows=1, cols=1)
log_table.alignment = WD_TABLE_ALIGNMENT.CENTER
c_log = log_table.cell(0, 0)
set_cell_background(c_log, '0F172A')
p_log = c_log.paragraphs[0]
p_log.paragraph_format.space_before = Pt(6)
p_log.paragraph_format.space_after = Pt(6)

with open('samples/live_demonstration_telemetry.txt', 'r', encoding='utf-8') as f:
    sample_text = f.read()

r_log = p_log.add_run(sample_text)
r_log.font.name = 'Consolas'
r_log.font.size = Pt(7.5)
r_log.font.color.rgb = RGBColor(226, 232, 240)

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# ----------------- SECTION 6: HARDWARE PINOUT & WIRING -----------------
add_header_p('6. Hardware Pinout & Wiring Specifications', level=1)

wire_tbl = doc.add_table(rows=10, cols=4)
wire_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
wire_headers = ['Device', 'Device Pin', 'Connected To', 'Functional Description']
for j, h in enumerate(wire_headers):
    c = wire_tbl.cell(0, j)
    set_cell_background(c, '1E3A8A')
    p = c.paragraphs[0]
    r = p.add_run(h)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    r.font.size = Pt(9.5)

wire_rows = [
    ('L298N Motor Driver', 'ENA', 'Arduino Pin 5 (PWM)', 'Left Motors Speed Control (PWM 0-255)'),
    ('L298N Motor Driver', 'ENB', 'Arduino Pin 6 (PWM)', 'Right Motors Speed Control (PWM 0-255)'),
    ('L298N Motor Driver', 'IN1, IN2', 'Arduino Pins 7, 8', 'Left Motors Direction Logic (H-Bridge)'),
    ('L298N Motor Driver', 'IN3, IN4', 'Arduino Pins 9, 11', 'Right Motors Direction Logic (H-Bridge)'),
    ('HC-SR04 Ultrasonic', 'TRIG', 'Arduino Pin 12', '10µs Ultrasonic Pulse Trigger Output'),
    ('HC-SR04 Ultrasonic', 'ECHO', 'Arduino Pin 13', 'Pulse Width Input for Echo Time Measurement'),
    ('Green Status LED', 'Anode (+)', 'Raspberry Pi BCM 27 (Pin 13)', 'WASTE Classification Indicator (Active HIGH)'),
    ('Red Status LED', 'Anode (+)', 'Raspberry Pi BCM 17 (Pin 11)', 'NON-WASTE Classification Indicator (Active HIGH)'),
    ('Power Supply', '7.4V - 12V Li-ion', 'L298N Power Terminals & Arduino GND', 'Dedicated High-Current Motor Power Rail (Common GND)')
]
for i, row_data in enumerate(wire_rows):
    row = wire_tbl.rows[i+1]
    bg = 'FFFFFF' if i % 2 == 0 else 'F8FAFC'
    for j, val in enumerate(row_data):
        c = row.cells[j]
        set_cell_background(c, bg)
        p = c.paragraphs[0]
        r = p.add_run(val)
        r.font.size = Pt(9)
        r.font.color.rgb = SLATE
        if j == 0:
            r.font.bold = True

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# ----------------- SECTION 7: WHAT TO INSTALL & HOW TO RUN -----------------
add_header_p('7. Installation & Execution Guide', level=1)

p_inst = doc.add_paragraph()
p_inst.add_run('A. Environment Setup & Dependencies:\n').font.bold = True
p_inst.add_run(
    '1. Clone the repository and navigate into the project directory:\n'
    '   git clone https://github.com/VAMSI-KRISHNA-N/cog.git\n'
    '   cd cog\n\n'
    '2. Create and activate a Python virtual environment:\n'
    '   python3 -m venv .venv\n'
    '   source .venv/bin/activate   (On Linux/Raspberry Pi)\n'
    '   .\\.venv\\Scripts\\activate      (On Windows)\n\n'
    '3. Install core dependencies:\n'
    '   pip install -r requirements.txt\n'
    '   (Dependencies: ultralytics>=8.0.0, opencv-python>=4.8.0, pyserial>=3.5, numpy>=1.23.0, RPi.GPIO on Linux)\n\n'
    '4. Verify Raspberry Pi hardware and environment readiness:\n'
    '   python3 check_pi.py\n'
)
p_inst.runs[0].font.size = Pt(10)
p_inst.runs[1].font.size = Pt(9)
p_inst.runs[1].font.name = 'Consolas'

p_run = doc.add_paragraph()
p_run.add_run('B. Execution Modes:\n').font.bold = True
p_run.add_run(
    '1. Physical Robot Execution on Raspberry Pi (Autonomous):\n'
    '   source .venv/bin/activate\n'
    '   python3 pi/robot.py --preview --no-shutdown\n\n'
    '2. Motor Diagnostic Verification:\n'
    '   python3 test_motors.py --port /dev/ttyACM0\n\n'
    '3. PC Simulation Mode (Interactive GUI with Mock Hardware & Synthetic Camera):\n'
    '   python simulate.py\n\n'
    '4. Automated Headless Systemd Service:\n'
    '   sudo cp systemd/cog-robot.service /etc/systemd/system/\n'
    '   sudo systemctl enable cog-robot.service\n'
)
p_run.runs[0].font.size = Pt(10)
p_run.runs[1].font.size = Pt(9)
p_run.runs[1].font.name = 'Consolas'

# ----------------- SECTION 8: CONCLUSION -----------------
add_header_p('8. Summary & Conclusion', level=1)
p_conc = doc.add_paragraph(
    'The COG Autonomous Waste-Detection Robot successfully meets all project requirements. The dual-tier architecture effectively '
    'combines the high-level perception capabilities of YOLOv8 running on the Raspberry Pi 4B with deterministic, safe low-level motor actuation '
    'managed by the Arduino Uno. Through the verified 7-state finite state machine, closed-loop visual servoing, and multi-frame consensus filtering, '
    'the robot achieved full autonomous floor navigation, target approach, and accurate waste classification as confirmed during the live faculty presentation.'
)
p_conc.runs[0].font.size = Pt(10)

doc.save('DA3_Submission_Report.docx')
print('DA3_Submission_Report.docx created successfully!')
